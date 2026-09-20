from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether


OUT = Path("artifacts")
OUT.mkdir(parents=True, exist_ok=True)

DOCX_NAME = "SABR_Claim_Status_Assistant_Technical_Case_Karimov_Muhamedzhan.docx"
PDF_NAME = "SABR_Claim_Status_Assistant_Technical_Case_Karimov_Muhamedzhan.pdf"
XLSX_NAME = "SABR_Claim_Status_Assistant_Model_Karimov_Muhamedzhan.xlsx"
PPTX_NAME = "SABR_Claim_Status_Assistant_Presentation_Karimov_Muhamedzhan.pptx"
ZIP_NAME = "SABR_Claim_Status_Assistant_Technical_Case_Package.zip"

PURPLE = "3B277A"
PURPLE_LIGHT = "EDE9FE"
GREEN = "1B8F5A"
GRAY = "5B6472"
LIGHT = "F6F7F9"
DARK = "16181D"
WHITE = "FFFFFF"
ORANGE = "D97706"

FIGMA_URL = "https://www.figma.com/design/eRYcFzUO9slgs4KrHdICbE"
REPO_URL = "https://github.com/memodlike/claim-status-assistant-technical-case"

SOURCES = [
    ("Sinoasia B&R", "Официальный сайт и текущий цифровой контекст", "https://www.sinoasia.kz/"),
    ("2GIS", "Публичные отзывы: обратная связь, процесс после ДТП, документы; также есть положительные отзывы", "https://2gis.kz/almaty/firm/70000001035253024/tab/reviews"),
    ("Страховой омбудсман", "Контекст страховых споров и досудебной коммуникации", "https://www.fomb.kz/"),
    ("Paul, 2021", "Задержки и слабая коммуникация в claims-процессах создают нагрузку для клиента", "https://doi.org/10.26180/14151785.v2"),
    ("Elgargouh et al.", "Сложный claims-процесс связан с риском ухода клиента", "https://doi.org/10.1109/CIST49399.2021.9357231"),
]

ASSUMPTIONS = {
    "Клиентская база": 250000,
    "Доля claim-related взаимодействий в год": 0.03,
    "Доля с повторным вопросом": 0.65,
    "Повторных контактов на случай": 2.2,
    "Минут на один контакт": 12,
    "Стоимость часа сотрудника, ₸": 3500,
    "Эскалаций/жалоб в месяц": 25,
    "Стоимость одной эскалации, ₸": 15000,
    "Потерянных продлений в месяц": 30,
    "Маржа на одно продление, ₸": 12000,
    "Стоимость MVP, ₸": 760000,
    "Стоимость полной интеграции, ₸": 12000000,
    "Цель снижения повторных контактов": 0.20,
}

CASES_PER_MONTH = ASSUMPTIONS["Клиентская база"] * ASSUMPTIONS["Доля claim-related взаимодействий в год"] / 12
REPEAT_CONTACTS = CASES_PER_MONTH * ASSUMPTIONS["Доля с повторным вопросом"] * ASSUMPTIONS["Повторных контактов на случай"]
STAFF_HOURS = REPEAT_CONTACTS * ASSUMPTIONS["Минут на один контакт"] / 60
CONTACT_COST = STAFF_HOURS * ASSUMPTIONS["Стоимость часа сотрудника, ₸"]
ESCALATION_COST = ASSUMPTIONS["Эскалаций/жалоб в месяц"] * ASSUMPTIONS["Стоимость одной эскалации, ₸"]
RENEWAL_MARGIN = ASSUMPTIONS["Потерянных продлений в месяц"] * ASSUMPTIONS["Маржа на одно продление, ₸"]
OPPORTUNITY = CONTACT_COST + ESCALATION_COST + RENEWAL_MARGIN


def money(v: float) -> str:
    return f"{v:,.0f} ₸".replace(",", " ")


def pct(v: float) -> str:
    return f"{v:.0%}"


# ---------------- DOCX ----------------

def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=DARK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_paragraph()
    p.style = doc.styles[f"Heading {level}"]
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor.from_string(PURPLE)
    return p


def add_body(doc: Document, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    for r in p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def add_bullets(doc: Document, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(10)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor.from_string(PURPLE)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor.from_string(PURPLE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("CLAIM STATUS ASSISTANT")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string(PURPLE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Технический кейс Product Masters")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sinoasia B&R / BCC Insurance • КАСКО / страховой случай")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Автор: Karimov Muhamedzhan • 20.09.2026")

    doc.add_paragraph()
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    shade_cell(cell, PURPLE_LIGHT)
    set_cell_text(
        cell,
        "Суть: после подачи заявления клиенту не всегда понятны статус, недостающие документы и следующий шаг. "
        "MVP — один экран со статусом + чек-лист документов + понятный ответ, который подтверждает сотрудник.",
        bold=True, color=PURPLE, size=11,
    )

    add_heading(doc, "1. Что именно решаем")
    add_body(doc, "Взята одна узкая часть процесса — проверка комплектности документов и коммуникация после регистрации страхового случая.")
    add_bullets(doc, [
        "что уже принято;",
        "каких документов не хватает;",
        "на каком этапе находится дело;",
        "что клиенту нужно сделать дальше.",
    ])
    add_body(doc, "ИИ не принимает решение о выплате или отказе. Он помогает структурировать информацию и подготовить понятный текст.")

    add_heading(doc, "2. Почему это болевая точка")
    add_body(doc, "Публичные отзывы используются только как pain-scan: они показывают повторяющиеся темы, но не являются статистикой всей клиентской базы.")
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(["Сигнал", "Что это значит", "Ограничение"]):
        shade_cell(t.rows[0].cells[i], PURPLE)
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=WHITE)
    data = [
        ("Трудно получить обратную связь", "Статус ищут через звонки/переписку", "Не у всех клиентов"),
        ("Вопросы после ДТП", "Нужен понятный следующий шаг", "Отзывы эмоциональны"),
        ("Споры по документам/оценке", "Нужна прозрачность процесса", "Не автоматизируем решение"),
        ("Есть положительные отзывы", "Картина неоднородна", "Нельзя делать общий негативный вывод"),
    ]
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell_text(cells[i], v, size=8)

    add_heading(doc, "3. Узкое место в пути клиента")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    headers = ["Этап", "Что происходит", "Боль"]
    for i, h in enumerate(headers):
        shade_cell(t.rows[0].cells[i], PURPLE)
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=WHITE)
    journey = [
        ("Страховой случай", "Клиент собирает материалы", "Не всегда понятно, что нужно именно для его ситуации"),
        ("Регистрация", "Заявление принимается", "Нет одного простого статуса"),
        ("Проверка комплекта", "Документы сверяются", "Недостающее выясняется через повторные контакты"),
        ("Оценка", "Материалы передаются дальше", "Непонятно, почему процесс стоит"),
        ("Решение", "Выплата / отказ / запрос", "Непрозрачность усиливает конфликт"),
    ]
    for row in journey:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell_text(cells[i], v, size=8)

    add_heading(doc, "4. Решение")
    add_body(doc, "Claim Status Assistant состоит из трёх частей:")
    add_bullets(doc, [
        "Клиентский экран: номер заявления, текущий статус, этапы, документы и следующий шаг.",
        "Рабочее место сотрудника: очередь заявлений, комплектность документов и черновик ответа.",
        "Экран эксперимента: primary metric, secondary metrics, guardrails и stop-критерии.",
    ])
    add_body(doc, f"Figma: {FIGMA_URL}")

    add_heading(doc, "5. Оценка проблемы в деньгах")
    add_body(doc, "Это модель для приоритизации. Неизвестные внутренние показатели обозначены как допущения.")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["Параметр", "Значение", "Статус"]):
        shade_cell(t.rows[0].cells[i], PURPLE)
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=WHITE)
    rows = [
        ("Клиентская база", "250 000", "рабочий масштаб"),
        ("Claim-related взаимодействия", "3% в год", "допущение"),
        ("Случаев в месяц", f"{CASES_PER_MONTH:.0f}", "расчёт"),
        ("Повторный вопрос", "65%", "допущение"),
        ("Повторных контактов", "2,2", "допущение"),
        ("Время контакта", "12 минут", "допущение"),
        ("Стоимость часа", "3 500 ₸", "допущение"),
    ]
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell_text(cells[i], v, size=8)

    add_body(doc, f"Повторные контакты: около {REPEAT_CONTACTS:.0f} в месяц, или {STAFF_HOURS:.0f} часов.")
    add_body(doc, f"Операционная стоимость повторных контактов: {money(CONTACT_COST)} в месяц.")
    add_body(doc, f"Эскалации/жалобы в модели: {money(ESCALATION_COST)} в месяц.")
    add_body(doc, f"Потеря маржи из-за непродлений в модели: {money(RENEWAL_MARGIN)} в месяц.")
    add_body(doc, f"Итого зона возможности: около {money(OPPORTUNITY)} в месяц.", bold_prefix="Итого зона возможности:")

    add_heading(doc, "6. Два варианта")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["", "MVP", "Полная интеграция"]):
        shade_cell(t.rows[0].cells[i], PURPLE)
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=WHITE)
    compare = [
        ("Срок", "≈ 1 неделя", "≈ 2 месяца"),
        ("Что внутри", "Статус, чек-лист, текст сотруднику", "АИС/CRM, кабинет, загрузка, уведомления"),
        ("Стоимость модели", "760 000 ₸", "12 000 000 ₸"),
        ("Риск", "Низкий", "Выше"),
        ("Решение", "Проверить ценность", "Только после подтверждения MVP"),
    ]
    for row in compare:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell_text(cells[i], v, size=8)

    add_heading(doc, "7. A/B-тест на 2 недели")
    add_body(doc, "Гипотеза: если статус, недостающие документы и следующий шаг видны в одном месте, повторные обращения снизятся минимум на 20%.")
    add_bullets(doc, [
        "50% — контроль: текущий процесс.",
        "50% — тест: Claim Status Assistant.",
        "Primary metric: повторные контакты на одно заявление за 14 дней.",
        "Secondary: время до полного пакета, доля полного пакета за 3 дня, CSAT.",
        "Guardrails: неверный статус, жалобы, ручная нагрузка, юридически рискованный текст.",
    ])
    add_body(doc, "Stop-критерии: неверных статусов >2%; жалобы в тесте выше контроля на 10%+; обработка сотрудником заметно замедлилась.")

    add_heading(doc, "8. Что сознательно не автоматизируем")
    add_bullets(doc, [
        "решение о выплате или отказе;",
        "методику оценки ущерба;",
        "обещание даты выплаты при неполном пакете;",
        "отправку AI-текста без проверки сотрудником;",
        "тяжёлую интеграцию до подтверждения эффекта MVP.",
    ])

    add_heading(doc, "9. Как использовался ИИ")
    add_body(doc, "ИИ применялся для поиска сигналов, кластеризации, расчётной модели и подготовки вариантов решения. После каждого шага результат ограничивался фактами и явными допущениями.")
    add_bullets(doc, [
        "Широкий список болей сужен до одного проверяемого сценария.",
        "Отзывы не используются как репрезентативная статистика.",
        "Все внутренние цифры без подтверждения обозначены как допущения.",
        "Из набора функций оставлены только MVP и следующий уровень интеграции.",
        "В эксперимент добавлены guardrails и stop-критерии.",
    ])

    add_heading(doc, "10. Источники")
    for name, note, url in SOURCES:
        add_body(doc, f"{name} — {note}. {url}")
    add_body(doc, f"Репозиторий: {REPO_URL}")

    doc.save(OUT / DOCX_NAME)


# ---------------- XLSX ----------------

def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"

    headers = ["Параметр", "Значение", "Комментарий"]
    ws.append(headers)
    assumptions_rows = [
        ("Клиентская база", 250000, "Рабочий масштаб модели"),
        ("Доля claim-related взаимодействий в год", 0.03, "Допущение"),
        ("Доля с повторным вопросом", 0.65, "Допущение"),
        ("Повторных контактов на случай", 2.2, "Допущение"),
        ("Минут на контакт", 12, "Допущение"),
        ("Стоимость часа сотрудника, ₸", 3500, "Допущение"),
        ("Эскалаций/жалоб в месяц", 25, "Допущение"),
        ("Стоимость одной эскалации, ₸", 15000, "Допущение"),
        ("Потерянных продлений в месяц", 30, "Допущение"),
        ("Маржа на продление, ₸", 12000, "Допущение"),
        ("Стоимость MVP, ₸", 760000, "Оценка"),
        ("Стоимость полной интеграции, ₸", 12000000, "Оценка"),
        ("Цель снижения повторных контактов", 0.20, "Гипотеза A/B"),
    ]
    for row in assumptions_rows:
        ws.append(row)

    econ = wb.create_sheet("Economics")
    econ.append(["Метрика", "Значение", "Формула / смысл"])
    econ_rows = [
        ("Случаев в месяц", "=Assumptions!B2*Assumptions!B3/12", "База × годовая доля / 12"),
        ("Повторных контактов в месяц", "=B2*Assumptions!B4*Assumptions!B5", "Случаи × доля × контакты"),
        ("Часов сотрудников", "=B3*Assumptions!B6/60", "Контакты × минуты / 60"),
        ("Стоимость повторных контактов, ₸", "=B4*Assumptions!B7", "Часы × стоимость часа"),
        ("Стоимость эскалаций, ₸", "=Assumptions!B8*Assumptions!B9", "Количество × стоимость"),
        ("Потеря маржи, ₸", "=Assumptions!B10*Assumptions!B11", "Непродления × маржа"),
        ("Зона возможности, ₸/мес", "=SUM(B5:B7)", "Операции + эскалации + маржа"),
        ("Экономия на повторных контактах при цели, ₸", "=B5*Assumptions!B14", "Консервативно: только повторные контакты"),
        ("Простой payback MVP, мес", "=Assumptions!B12/B9", "MVP / консервативная экономия"),
    ]
    for row in econ_rows:
        econ.append(row)

    exp = wb.create_sheet("Experiment")
    exp.append(["Блок", "Показатель", "Критерий"])
    experiment_rows = [
        ("Primary", "Повторные контакты на заявление за 14 дней", "−20% или лучше"),
        ("Secondary", "Время до полного пакета", "−15% или лучше"),
        ("Secondary", "Полный пакет за 3 дня", "Рост"),
        ("Secondary", "CSAT", "Не снижается"),
        ("Guardrail", "Неверный статус", "≤2%"),
        ("Guardrail", "Жалобы", "Не выше контроля на 10%+"),
        ("Guardrail", "Ручная нагрузка", "Не растёт заметно"),
    ]
    for row in experiment_rows:
        exp.append(row)

    src = wb.create_sheet("Sources")
    src.append(["Источник", "Что использовано", "URL"])
    for name, note, url in SOURCES:
        src.append([name, note, url])

    fill = PatternFill("solid", fgColor=PURPLE)
    fill2 = PatternFill("solid", fgColor=PURPLE_LIGHT)
    white_font = Font(color=WHITE, bold=True)
    title_font = Font(color=PURPLE, bold=True)
    thin = Side(style="thin", color="D9DCE1")

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, sheet.max_column + 1):
            width = max(14, min(48, max(len(str(sheet.cell(r, col).value or "")) for r in range(1, sheet.max_row + 1)) + 2))
            sheet.column_dimensions[get_column_letter(col)].width = width

    for row in [3, 4, 14]:
        ws.cell(row, 2).number_format = "0%"
    for row in range(2, econ.max_row + 1):
        if "₸" in str(econ.cell(row, 1).value):
            econ.cell(row, 2).number_format = '#,##0 "₸"'
    econ["B10"].number_format = "0.0"

    wb.save(OUT / XLSX_NAME)


# ---------------- PPTX ----------------

def add_text(slide, x, y, w, h, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = PPt(size)
    r.font.bold = bold
    r.font.color.rgb = PRGBColor.from_string(color)
    return box


def add_rect(slide, x, y, w, h, fill=WHITE, line="E5E7EB", radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, PInches(x), PInches(y), PInches(w), PInches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = PRGBColor.from_string(fill)
    s.line.color.rgb = PRGBColor.from_string(line)
    return s


def add_slide_title(slide, title, subtitle=None):
    add_text(slide, 0.65, 0.35, 12.0, 0.5, title, 25, True, PURPLE)
    if subtitle:
        add_text(slide, 0.65, 0.88, 12.0, 0.42, subtitle, 11, False, GRAY)


def build_pptx():
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]

    # 1 cover
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, 13.333, 7.5, PURPLE, PURPLE)
    add_text(s, 0.8, 1.35, 11.8, 1.0, "CLAIM STATUS ASSISTANT", 34, True, WHITE)
    add_text(s, 0.8, 2.45, 10.8, 0.7, "Технический кейс Product Masters", 20, False, "E9E4FA")
    add_text(s, 0.8, 3.3, 10.8, 0.8, "Sinoasia B&R / BCC Insurance\nКАСКО / страховой случай", 16, False, WHITE)
    add_text(s, 0.8, 6.3, 11.5, 0.45, "Karimov Muhamedzhan • 20.09.2026", 11, False, "D9D4EB")

    # 2 problem
    s = prs.slides.add_slide(blank)
    add_slide_title(s, "Проблема", "После подачи заявления клиенту не всегда понятно, что происходит дальше.")
    items = [
        ("01", "Что уже принято"),
        ("02", "Каких документов не хватает"),
        ("03", "На каком этапе находится дело"),
        ("04", "Какой следующий шаг"),
    ]
    for i, (n, t) in enumerate(items):
        x = 0.75 + (i % 2) * 6.1
        y = 1.65 + (i // 2) * 2.15
        add_rect(s, x, y, 5.55, 1.55, LIGHT, "E7E8EC", True)
        add_text(s, x + 0.25, y + 0.2, 0.65, 0.4, n, 16, True, GREEN)
        add_text(s, x + 1.05, y + 0.2, 4.1, 0.85, t, 20, True, DARK)
    add_text(s, 0.75, 6.3, 11.8, 0.45, "Выбран один узкий сценарий, а не автоматизация всего урегулирования.", 12, True, PURPLE)

    # 3 evidence
    s = prs.slides.add_slide(blank)
    add_slide_title(s, "Что подтверждает боль", "Отзывы — сигнал, а не статистика всей клиентской базы.")
    cards = [
        ("Обратная связь", "В публичных отзывах встречаются сложности с дозвоном и получением ответа."),
        ("Процесс после ДТП", "Клиенту нужен понятный статус и следующий шаг."),
        ("Документы и оценка", "Непрозрачность усиливает конфликт вокруг решения."),
        ("Контрсигнал", "Есть и положительные отзывы — картина неоднородна."),
    ]
    for i, (h, body) in enumerate(cards):
        x = 0.65 + i * 3.12
        add_rect(s, x, 1.7, 2.8, 3.65, WHITE, "DDD7EF", True)
        add_text(s, x + 0.22, 1.95, 2.35, 0.55, h, 16, True, PURPLE)
        add_text(s, x + 0.22, 2.7, 2.35, 1.8, body, 12, False, DARK)
    add_text(s, 0.75, 6.05, 11.7, 0.8, "Исследования claims-experience также связывают задержки, сложность процесса и слабую коммуникацию с ухудшением клиентского опыта.", 12, False, GRAY)

    # 4 bottleneck
    s = prs.slides.add_slide(blank)
    add_slide_title(s, "Узкое место", "Проверка комплектности и коммуникация после регистрации.")
    stages = ["Страховой\nслучай", "Регистрация", "Проверка\nкомплекта", "Оценка", "Решение"]
    for i, st in enumerate(stages):
        x = 0.65 + i * 2.48
        fill = PURPLE_LIGHT if i == 2 else LIGHT
        line = PURPLE if i == 2 else "E2E4E8"
        add_rect(s, x, 2.25, 2.0, 1.15, fill, line, True)
        add_text(s, x + 0.1, 2.52, 1.8, 0.6, st, 14, True, PURPLE if i == 2 else DARK, PP_ALIGN.CENTER)
        if i < 4:
            add_text(s, x + 2.0, 2.56, 0.45, 0.35, "→", 18, True, GRAY, PP_ALIGN.CENTER)
    add_rect(s, 3.85, 4.15, 5.55, 1.35, "FFF7ED", "F6C98A", True)
    add_text(s, 4.1, 4.42, 5.0, 0.8, "Здесь чаще всего нужен повторный контакт:\n«что не хватает и что дальше?»", 17, True, ORANGE, PP_ALIGN.CENTER)

    # 5 solution
    s = prs.slides.add_slide(blank)
    add_slide_title(s, "MVP: один понятный статус", "ИИ помогает подготовить информацию, но сотрудник остаётся в контуре.")
    cards = [
        ("Клиент", ["статус", "этапы", "недостающие документы", "следующий шаг"]),
        ("Сотрудник", ["очередь заявлений", "комплектность", "проблемные позиции", "подтверждение ответа"]),
        ("Эксперимент", ["primary metric", "secondary metrics", "guardrails", "stop criteria"]),
    ]
    for i, (h, bullets) in enumerate(cards):
        x = 0.65 + i * 4.15
        add_rect(s, x, 1.65, 3.75, 4.55, WHITE, "DDD7EF", True)
        add_text(s, x + 0.25, 1.93, 3.2, 0.45, h, 18, True, PURPLE)
        y = 2.7
        for b in bullets:
            add_text(s, x + 0.3, y, 3.1, 0.5, "• " + b, 12, False, DARK)
            y += 0.65
    add_text(s, 0.75, 6.55, 11.8, 0.35, "Figma: " + FIGMA_URL, 9, False, GRAY)

    # 6 economics
    s = prs.slides.add_slide(blank)
    add_slide_title(s, "Экономика: порядок величины", "Все внутренние цифры ниже — прозрачные допущения модели.")
    kpis = [
        ("625", "случаев / мес."),
        (f"{REPEAT_CONTACTS:.0f}", "повторных контактов / мес."),
        (money(CONTACT_COST), "стоимость повторных контактов"),
        (money(OPPORTUNITY), "общая зона возможности / мес."),
    ]
    for i, (v, lab) in enumerate(kpis):
        x = 0.65 + i * 3.15
        add_rect(s, x, 1.65, 2.85, 1.65, LIGHT, "E0E2E6", True)
        add_text(s, x + 0.2, 1.95, 2.45, 0.55, v, 22, True, PURPLE, PP_ALIGN.CENTER)
        add_text(s, x + 0.2, 2.58, 2.45, 0.42, lab, 10, False, GRAY, PP_ALIGN.CENTER)
    add_rect(s, 0.65, 4.0, 5.85, 1.7, PURPLE_LIGHT, PURPLE_LIGHT, True)
    add_text(s, 0.95, 4.3, 5.2, 0.45, "MVP: ≈ 760 000 ₸ • ≈ 1 неделя", 17, True, PURPLE)
    add_text(s, 0.95, 4.9, 5.2, 0.45, "Сначала проверяем ценность дешёво.", 12, False, DARK)
    add_rect(s, 6.85, 4.0, 5.85, 1.7, LIGHT, "E0E2E6", True)
    add_text(s, 7.15, 4.3, 5.2, 0.45, "Интеграция: ≈ 12 млн ₸ • ≈ 2 месяца", 17, True, DARK)
    add_text(s, 7.15, 4.9, 5.2, 0.45, "Только если MVP подтвердит эффект.", 12, False, GRAY)

    # 7 experiment
    s = prs.slides.add_slide(blank)
    add_slide_title(s, "A/B-тест на 2 недели", "Проверяем поведение, а не клики по интерфейсу.")
    add_rect(s, 0.75, 1.55, 5.75, 1.75, LIGHT, "E1E3E7", True)
    add_text(s, 1.0, 1.85, 5.2, 0.45, "CONTROL", 13, True, GRAY)
    add_text(s, 1.0, 2.35, 5.1, 0.65, "Текущий процесс и текущие каналы уточнения статуса.", 15, True, DARK)
    add_rect(s, 6.85, 1.55, 5.75, 1.75, PURPLE_LIGHT, PURPLE, True)
    add_text(s, 7.1, 1.85, 5.2, 0.45, "TEST", 13, True, PURPLE)
    add_text(s, 7.1, 2.35, 5.1, 0.65, "Статус + чек-лист + подтверждённый сотрудником ответ.", 15, True, PURPLE)

    metrics = [
        ("Primary", "Повторные контакты", "−20%"),
        ("Secondary", "Время до полного пакета", "−15%"),
        ("Guardrail", "Неверный статус", "≤2%"),
        ("Stop", "Жалобы vs control", "не +10%"),
    ]
    for i, (k, m, v) in enumerate(metrics):
        x = 0.75 + i * 3.05
        add_rect(s, x, 4.0, 2.75, 1.65, WHITE, "E0E2E6", True)
        add_text(s, x + 0.18, 4.2, 2.4, 0.3, k, 10, True, GREEN)
        add_text(s, x + 0.18, 4.65, 2.4, 0.5, m, 12, True, DARK)
        add_text(s, x + 0.18, 5.18, 2.4, 0.35, v, 16, True, PURPLE)

    # 8 decision
    s = prs.slides.add_slide(blank)
    add_slide_title(s, "Решение по результату", "Не строить тяжёлую интеграцию до доказанного эффекта.")
    add_rect(s, 0.75, 1.65, 11.85, 1.65, PURPLE, PURPLE, True)
    add_text(s, 1.05, 2.05, 11.2, 0.85, "Если MVP снижает повторные контакты без роста жалоб — переходить к интеграции с АИС/CRM.", 22, True, WHITE, PP_ALIGN.CENTER)
    add_text(s, 0.85, 4.2, 11.6, 0.5, "Что не автоматизируем", 18, True, PURPLE, PP_ALIGN.CENTER)
    add_text(s, 1.1, 4.95, 11.1, 0.95, "Решение о выплате/отказе • оценку ущерба • обещание даты выплаты • отправку AI-текста без проверки", 15, False, DARK, PP_ALIGN.CENTER)
    add_text(s, 0.9, 6.35, 11.5, 0.3, REPO_URL, 9, False, GRAY, PP_ALIGN.CENTER)

    prs.save(OUT / PPTX_NAME)


# ---------------- PDF ----------------

def build_pdf():
    regular = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    pdfmetrics.registerFont(TTFont("DejaVu", regular))
    pdfmetrics.registerFont(TTFont("DejaVuBold", bold))

    styles = getSampleStyleSheet()
    title = ParagraphStyle("CaseTitle", parent=styles["Title"], fontName="DejaVuBold", fontSize=22, leading=26, textColor=colors.HexColor("#" + PURPLE), alignment=TA_CENTER, spaceAfter=10)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName="DejaVuBold", fontSize=14, leading=18, textColor=colors.HexColor("#" + PURPLE), spaceBefore=8, spaceAfter=6)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="DejaVu", fontSize=9.5, leading=14, textColor=colors.HexColor("#" + DARK), spaceAfter=5)
    small = ParagraphStyle("Small", parent=body, fontSize=8, leading=11, textColor=colors.HexColor("#" + GRAY))
    callout = ParagraphStyle("Callout", parent=body, fontName="DejaVuBold", fontSize=11, leading=16, textColor=colors.HexColor("#" + PURPLE), backColor=colors.HexColor("#" + PURPLE_LIGHT), borderPadding=10)

    doc = SimpleDocTemplate(str(OUT / PDF_NAME), pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=14*mm, bottomMargin=14*mm)
    story = []
    story += [
        Paragraph("CLAIM STATUS ASSISTANT", title),
        Paragraph("Технический кейс Product Masters", ParagraphStyle("Sub", parent=body, alignment=TA_CENTER, fontSize=12, textColor=colors.HexColor("#"+GRAY))),
        Paragraph("Sinoasia B&R / BCC Insurance • КАСКО / страховой случай", ParagraphStyle("Sub2", parent=body, alignment=TA_CENTER, fontSize=9, textColor=colors.HexColor("#"+GREEN))),
        Spacer(1, 8),
        Paragraph("После подачи заявления клиенту не всегда понятны статус, недостающие документы и следующий шаг. MVP — один экран со статусом, чек-листом и понятным ответом, который подтверждает сотрудник.", callout),
        Spacer(1, 8),
    ]

    sections = [
        ("1. Что решаем", [
            "Выбран один узкий участок: проверка комплектности документов и коммуникация после регистрации страхового случая.",
            "Клиенту нужно быстро понять: что принято, чего не хватает, где находится дело и что делать дальше.",
            "<b>ИИ не принимает решение о выплате или отказе.</b>",
        ]),
        ("2. Сигналы боли", [
            "Публичные отзывы дают повторяющиеся сигналы по обратной связи, процессу после ДТП и вопросам по документам. При этом есть и положительные отзывы.",
            "Поэтому отзывы используются как pain-scan, а не как статистика всей клиентской базы.",
            "Исследования claims-experience также связывают задержки, сложность процесса и слабую коммуникацию с ухудшением клиентского опыта.",
        ]),
        ("3. Решение", [
            "Клиентский экран: номер заявления, текущий статус, этапы, недостающие документы и следующий шаг.",
            "Рабочее место сотрудника: очередь заявлений, комплектность, проблемные позиции и черновик ответа.",
            "Экран эксперимента: primary metric, secondary metrics, guardrails и stop-критерии.",
            f"Figma: {FIGMA_URL}",
        ]),
        ("4. Экономика", [
            f"При учебных допущениях получается около {CASES_PER_MONTH:.0f} claim-related случаев и {REPEAT_CONTACTS:.0f} повторных контактов в месяц.",
            f"Операционная стоимость повторных контактов — около <b>{money(CONTACT_COST)}</b> в месяц.",
            f"С учётом модели эскалаций и непродлений общая зона возможности — около <b>{money(OPPORTUNITY)}</b> в месяц.",
            "Все неизвестные внутренние показатели — допущения. Перед реальным решением их нужно заменить данными контакт-центра, выплат и продлений.",
        ]),
        ("5. Два варианта", [
            "<b>MVP:</b> ≈1 неделя, статус + чек-лист + текст сотруднику. Оценка стоимости — 760 000 ₸.",
            "<b>Полная интеграция:</b> ≈2 месяца, АИС/CRM + кабинет + загрузка документов + уведомления. Оценка — 12 млн ₸.",
            "Логика: сначала подтвердить эффект дешёвым MVP, затем инвестировать в интеграцию.",
        ]),
        ("6. A/B-тест", [
            "50% — контроль, 50% — тест, срок — 2 недели.",
            "<b>Primary metric:</b> повторные контакты на одно заявление за 14 дней. Цель — −20% или лучше.",
            "<b>Secondary:</b> время до полного пакета, полный пакет за 3 дня, CSAT.",
            "<b>Guardrails:</b> неверный статус, жалобы, ручная нагрузка, юридически рискованный текст.",
            "Stop-критерии: неверных статусов >2%; жалобы выше контроля на 10%+; сотрудникам стало заметно дольше.",
        ]),
        ("7. Ограничения", [
            "Не автоматизируем выплату/отказ, методику оценки ущерба и обещание даты выплаты.",
            "Не отправляем AI-текст без проверки сотрудником.",
            "Отзывы не трактуем как репрезентативную статистику.",
            "Финансовая модель показывает порядок величины, а не подтверждённый эффект.",
        ]),
    ]

    for head, paras in sections:
        story.append(Paragraph(head, h1))
        for p in paras:
            story.append(Paragraph(p, body))

    story.append(PageBreak())
    story.append(Paragraph("Источники", h1))
    src_data = [["Источник", "Что использовано", "URL"]]
    for n, note, url in SOURCES:
        src_data.append([Paragraph(n, small), Paragraph(note, small), Paragraph(url, small)])
    table = Table(src_data, colWidths=[35*mm, 80*mm, 65*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#"+PURPLE)),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "DejaVuBold"),
        ("FONTNAME", (0,1), (-1,-1), "DejaVu"),
        ("FONTSIZE", (0,0), (-1,-1), 7.5),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#D9DCE1")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F7F7F9")]),
        ("LEFTPADDING", (0,0), (-1,-1), 5),
        ("RIGHTPADDING", (0,0), (-1,-1), 5),
        ("TOPPADDING", (0,0), (-1,-1), 5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story.append(table)
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"Репозиторий: {REPO_URL}", small))
    doc.build(story)


def build_readme_for_typeform():
    text = f"""Claim Status Assistant — технический кейс Product Masters

Суть:
после подачи заявления клиенту не всегда понятны статус, недостающие документы и следующий шаг.

Решение:
экран статуса + чек-лист документов + понятный ответ, который подтверждает сотрудник.

Проверка:
A/B-тест на 2 недели. Главная цель — снизить повторные обращения минимум на 20% без роста жалоб и ошибок статуса.

Важно:
публичные отзывы используются как сигнал боли, а неизвестные внутренние цифры в модели обозначены как допущения.

Figma:
{FIGMA_URL}

GitHub:
{REPO_URL}
"""
    (OUT / "README_for_Typeform.txt").write_text(text, encoding="utf-8")


def build_zip():
    files = [DOCX_NAME, PDF_NAME, XLSX_NAME, PPTX_NAME, "README_for_Typeform.txt"]
    with ZipFile(OUT / ZIP_NAME, "w", ZIP_DEFLATED) as z:
        for name in files:
            z.write(OUT / name, arcname=name)


if __name__ == "__main__":
    build_docx()
    build_xlsx()
    build_pptx()
    build_pdf()
    build_readme_for_typeform()
    build_zip()
    print("Built:")
    for p in sorted(OUT.iterdir()):
        print(f" - {p.name}: {p.stat().st_size} bytes")
